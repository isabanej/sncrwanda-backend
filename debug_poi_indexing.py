import openpyxl

wb = openpyxl.load_workbook('frontend-new/public/SNC Rwanda - 12 Month Income & Expense Projection = CashflowProjection.xlsx')

# Check Cashflow Statement sheet
cashflow_sheet = wb['Cashflow Statement']
print("="* 80)
print("CASHFLOW STATEMENT SHEET - HEADER ROW")
print("=" * 80)
print("Row 1 (index 0 in POI):")
for col in range(0, 15):
    cell = cashflow_sheet.cell(1, col + 1)
    if cell.value:
        print(f"  Col {col} (POI) = Col {col+1} (Excel): '{cell.value}'")

# Check Revenue sheet
revenue_sheet = wb['School Fees received - Revenue']
print("\n" + "=" * 80)
print("REVENUE SHEET - HEADER AND DATA")
print("=" * 80)

print("\nRow 3 (Header):")
for col in range(0, 15):
    cell = revenue_sheet.cell(3, col + 1)
    if cell.value:
        print(f"  Col {col} (POI) = Col {col+1} (Excel): '{cell.value}'")

print("\nData Rows (starting from row 4):")
for row in range(4, 12):
    name_cell = revenue_sheet.cell(row, 1)
    if name_cell.value:
        print(f"\nRow {row} (POI row {row-1}): Student='{name_cell.value}'")
        # Check all columns for this row
        for col in range(2, 15):
            cell = revenue_sheet.cell(row, col + 1)
            if cell.value and isinstance(cell.value, (int, float)) and cell.value > 0:
                header_cell = revenue_sheet.cell(3, col + 1)
                print(f"  Col {col} ({header_cell.value if header_cell.value else 'N/A'}): ${cell.value:,.2f}")

wb.close()
