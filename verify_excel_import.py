#!/usr/bin/env python3
"""
Comprehensive Excel Import Verification Script
This script:
1. Imports historical cashflow data from Excel
2. Verifies that database data matches Excel data exactly
3. Generates a detailed comparison report
"""

import requests
import openpyxl
import psycopg2
from decimal import Decimal
from datetime import datetime
import sys

# Configuration
LEDGER_SERVICE_URL = "http://localhost:8082"
AUTH_SERVICE_URL = "http://localhost:9092"
EXCEL_FILE = "frontend-new/public/SNC Rwanda - 12 Month Income & Expense Projection = CashflowProjection.xlsx"

DB_CONFIG = {
    'host': 'localhost',
    'port': 5432,
    'database': 'sncrwanda',
    'user': 'postgres',
    'password': 'postgres'
}

# Default org and user IDs
ORG_ID = "00000000-0000-0000-0000-000000000001"
ADMIN_USER_ID = "00000000-0000-0000-0000-000000000001"

def login():
    """Login and get access token"""
    print("🔐 Logging in to auth service...")
    response = requests.post(
        f"{AUTH_SERVICE_URL}/api/auth/login",
        json={"username": "emino", "password": "123456"}
    )
    if response.status_code == 200:
        token = response.json().get("accessToken")
        print("✅ Login successful")
        return token
    else:
        print(f"❌ Login failed: {response.status_code}")
        print(response.text)
        sys.exit(1)

def parse_excel_data():
    """Parse Excel file and extract all data"""
    print(f"\n📊 Parsing Excel file: {EXCEL_FILE}")
    
    wb = openpyxl.load_workbook(EXCEL_FILE)
    cashflow_sheet = wb['Cashflow Statement']
    revenue_sheet = wb.get_sheet_by_name('School Fees received - Revenue')
    
    data = {
        'months': [],
        'fees': {},
        'expenses': {},
        'petty_cash': {}
    }
    
    # Parse month headers from row 1 (columns D onwards)
    header_row = cashflow_sheet[1]
    for col_idx in range(3, header_row[0].parent.max_column):  # Start from column D (index 3)
        cell = cashflow_sheet.cell(1, col_idx + 1)
        if cell.value:
            month_name = str(cell.value).strip()
            # Check if it's a month name
            month_names = ["January", "February", "March", "April", "May", "June",
                          "July", "August", "September", "October", "November", "December",
                          "Jan", "Feb", "Fab", "Mar", "Apr", "Jun", "Jul", "Aug", "Sept", "Sep", "Oct", "Nov", "Dec"]
            if any(month_name.lower() == m.lower() for m in month_names):
                data['months'].append((col_idx + 1, month_name))  # Store column index and name
    
    print(f"   Found {len(data['months'])} months: {[m[1] for m in data['months']]}")
    
    # Parse student fees from revenue sheet
    if revenue_sheet:
        print("   Parsing student fees...")
        for row_idx in range(2, revenue_sheet.max_row + 1):
            student_name_cell = revenue_sheet.cell(row_idx, 1)
            if not student_name_cell.value:
                continue
            
            student_name = str(student_name_cell.value).strip()
            if student_name.lower() == 'total':
                continue
            
            fee_type_cell = revenue_sheet.cell(row_idx, 2)
            fee_type = str(fee_type_cell.value).strip() if fee_type_cell.value else "School Fees"
            
            for col_idx, month_name in data['months']:
                amount_cell = revenue_sheet.cell(row_idx, col_idx)
                if amount_cell.value and isinstance(amount_cell.value, (int, float)) and amount_cell.value > 0:
                    if month_name not in data['fees']:
                        data['fees'][month_name] = []
                    data['fees'][month_name].append({
                        'student_name': student_name,
                        'fee_type': fee_type,
                        'amount': Decimal(str(amount_cell.value))
                    })
    
    # Parse expenses from cashflow sheet
    print("   Parsing expenses...")
    expense_categories = {}
    for row_idx in range(10, 30):  # Rows 10-29 typically contain expenses
        category_cell = cashflow_sheet.cell(row_idx, 2)  # Column B
        if not category_cell.value:
            continue
        
        category = str(category_cell.value).strip()
        
        # Filter out non-expense rows
        if (not category or 
            category.lower() in ['total', 'cash available', 'total expenses'] or
            'petty cash' in category.lower() or
            'Cash In' in category or
            'Cash Out' in category):
            continue
        
        expense_categories[row_idx] = category
        
        for col_idx, month_name in data['months']:
            amount_cell = cashflow_sheet.cell(row_idx, col_idx)
            if amount_cell.value and isinstance(amount_cell.value, (int, float)) and amount_cell.value > 0:
                if month_name not in data['expenses']:
                    data['expenses'][month_name] = {}
                if category not in data['expenses'][month_name]:
                    data['expenses'][month_name][category] = Decimal(0)
                data['expenses'][month_name][category] += Decimal(str(amount_cell.value))
    
    # Parse petty cash IN/OUT
    print("   Parsing petty cash transactions...")
    for row_idx in range(1, cashflow_sheet.max_row + 1):
        category_cell = cashflow_sheet.cell(row_idx, 2)  # Column B
        if not category_cell.value:
            continue
        
        category = str(category_cell.value).strip()
        
        if category.lower() == 'petty cash(in)':
            for col_idx, month_name in data['months']:
                amount_cell = cashflow_sheet.cell(row_idx, col_idx)
                if amount_cell.value and isinstance(amount_cell.value, (int, float)) and amount_cell.value > 0:
                    if month_name not in data['petty_cash']:
                        data['petty_cash'][month_name] = {'IN': Decimal(0), 'OUT': Decimal(0)}
                    data['petty_cash'][month_name]['IN'] += Decimal(str(amount_cell.value))
        
        elif category.lower() == 'petty cash(out)':
            for col_idx, month_name in data['months']:
                amount_cell = cashflow_sheet.cell(row_idx, col_idx)
                if amount_cell.value and isinstance(amount_cell.value, (int, float)) and amount_cell.value > 0:
                    if month_name not in data['petty_cash']:
                        data['petty_cash'][month_name] = {'IN': Decimal(0), 'OUT': Decimal(0)}
                    data['petty_cash'][month_name]['OUT'] += Decimal(str(amount_cell.value))
    
    wb.close()
    
    # Print summary
    print("\n📈 Excel Data Summary:")
    total_fees = sum(sum(f['amount'] for f in fees) for fees in data['fees'].values())
    total_expenses = sum(sum(amounts.values()) for amounts in data['expenses'].values())
    total_petty_in = sum(pc.get('IN', 0) for pc in data['petty_cash'].values())
    total_petty_out = sum(pc.get('OUT', 0) for pc in data['petty_cash'].values())
    
    print(f"   Total Student Fees: ${total_fees:,.2f}")
    print(f"   Total Expenses: ${total_expenses:,.2f}")
    print(f"   Total Petty Cash IN: ${total_petty_in:,.2f}")
    print(f"   Total Petty Cash OUT: ${total_petty_out:,.2f}")
    
    return data

def import_excel_data(token):
    """Import Excel file via API"""
    print("\n📤 Importing Excel file to ledger service...")
    
    with open(EXCEL_FILE, 'rb') as f:
        files = {'file': (EXCEL_FILE.split('/')[-1], f, 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}
        params = {
            'orgId': ORG_ID,
            'importedBy': ADMIN_USER_ID
        }
        
        response = requests.post(
            f"{LEDGER_SERVICE_URL}/api/cashflow/import/excel",
            files=files,
            params=params,
            headers={"Authorization": f"Bearer {token}"}
        )
        
        if response.status_code == 200:
            result = response.json()
            print("✅ Import successful!")
            print(f"   Periods created: {result.get('periodsCreated', 0)}")
            print(f"   Fees imported: {result.get('feesImported', 0)}")
            print(f"   Expenses imported: {result.get('expensesImported', 0)}")
            print(f"   Petty cash imported: {result.get('pettyCashImported', 0)}")
            return result
        else:
            print(f"❌ Import failed: {response.status_code}")
            print(response.text)
            return None

def verify_database_data(excel_data):
    """Verify database data matches Excel"""
    print("\n🔍 Verifying database data...")
    
    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()
    
    issues = []
    
    # Get current year
    current_year = datetime.now().year
    
    # Verify periods created
    print("   Checking periods...")
    cur.execute("""
        SELECT year, month, period_name 
        FROM ledger.cashflow_periods 
        WHERE org_id = %s
        ORDER BY year, month
    """, (ORG_ID,))
    periods = cur.fetchall()
    
    if len(periods) != len(excel_data['months']):
        issues.append(f"Period count mismatch: Expected {len(excel_data['months'])}, found {len(periods)}")
    
    # Verify student fees
    print("   Checking student fees...")
    for month_name, fees_list in excel_data['fees'].items():
        month_num = get_month_number(month_name)
        
        cur.execute("""
            SELECT COUNT(*), SUM(amount_paid)
            FROM ledger.student_fee_payments sfp
            JOIN ledger.cashflow_periods p ON sfp.period_id = p.id
            WHERE p.year = %s AND p.month = %s AND p.org_id = %s
        """, (current_year, month_num, ORG_ID))
        
        db_count, db_total = cur.fetchone()
        excel_count = len(fees_list)
        excel_total = sum(f['amount'] for f in fees_list)
        
        if db_count != excel_count:
            issues.append(f"{month_name} fees count: Expected {excel_count}, found {db_count}")
        
        if db_total is not None and abs(float(db_total) - float(excel_total)) > 0.01:
            issues.append(f"{month_name} fees total: Expected ${excel_total:,.2f}, found ${db_total:,.2f}")
    
    # Verify expenses
    print("   Checking expenses...")
    for month_name, expenses_dict in excel_data['expenses'].items():
        month_num = get_month_number(month_name)
        
        for category, excel_amount in expenses_dict.items():
            cur.execute("""
                SELECT SUM(amount)
                FROM ledger.cashflow_expenses e
                JOIN ledger.cashflow_periods p ON e.period_id = p.id
                WHERE p.year = %s AND p.month = %s AND p.org_id = %s AND e.category = %s
            """, (current_year, month_num, ORG_ID, category))
            
            result = cur.fetchone()
            db_amount = result[0] if result[0] is not None else Decimal(0)
            
            if abs(float(db_amount) - float(excel_amount)) > 0.01:
                issues.append(f"{month_name} {category}: Expected ${excel_amount:,.2f}, found ${db_amount:,.2f}")
    
    # Verify petty cash
    print("   Checking petty cash...")
    for month_name, petty_cash_dict in excel_data['petty_cash'].items():
        month_num = get_month_number(month_name)
        
        for tx_type, excel_amount in petty_cash_dict.items():
            if excel_amount == 0:
                continue
            
            cur.execute("""
                SELECT SUM(amount)
                FROM ledger.petty_cash_transactions pct
                JOIN ledger.cashflow_periods p ON pct.period_id = p.id
                WHERE p.year = %s AND p.month = %s AND p.org_id = %s AND pct.transaction_type = %s
            """, (current_year, month_num, ORG_ID, tx_type))
            
            result = cur.fetchone()
            db_amount = result[0] if result[0] is not None else Decimal(0)
            
            if abs(float(db_amount) - float(excel_amount)) > 0.01:
                issues.append(f"{month_name} Petty Cash {tx_type}: Expected ${excel_amount:,.2f}, found ${db_amount:,.2f}")
    
    cur.close()
    conn.close()
    
    return issues

def get_month_number(month_name):
    """Convert month name to number"""
    month_map = {
        'january': 1, 'jan': 1,
        'february': 2, 'feb': 2, 'fab': 2,
        'march': 3, 'mar': 3,
        'april': 4, 'apr': 4,
        'may': 5,
        'june': 6, 'jun': 6,
        'july': 7, 'jul': 7,
        'august': 8, 'aug': 8,
        'september': 9, 'sept': 9, 'sep': 9,
        'october': 10, 'oct': 10,
        'november': 11, 'nov': 11,
        'december': 12, 'dec': 12
    }
    return month_map.get(month_name.lower(), 1)

def main():
    print("=" * 80)
    print("  EXCEL IMPORT VERIFICATION TOOL")
    print("=" * 80)
    
    # Step 1: Parse Excel data
    excel_data = parse_excel_data()
    
    # Step 2: Login
    token = login()
    
    # Step 3: Import Excel data
    import_result = import_excel_data(token)
    if not import_result:
        print("\n❌ Import failed. Exiting.")
        sys.exit(1)
    
    # Step 4: Verify data
    issues = verify_database_data(excel_data)
    
    # Step 5: Report results
    print("\n" + "=" * 80)
    if not issues:
        print("✅ VERIFICATION COMPLETE: All data matches perfectly!")
    else:
        print(f"⚠️  VERIFICATION COMPLETE: Found {len(issues)} discrepancies:")
        for issue in issues:
            print(f"   • {issue}")
    print("=" * 80)

if __name__ == "__main__":
    main()
