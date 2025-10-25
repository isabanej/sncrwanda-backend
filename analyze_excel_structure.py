import openpyxl

# Load the Excel file
wb = openpyxl.load_workbook(r"c:\dev\sncrwanda-backend\frontend-new\public\SNC Rwanda - 12 Month Income & Expense Projection = CashflowProjection.xlsx")

# Check Revenue sheet structure
revenue_sheet = wb["School Fees received - Revenue"]

print("REVENUE SHEET STRUCTURE:")
print("=" * 60)

# Print header row (row 3)
header_row = 3
print(f"\nRow {header_row} (Header):")
for col_idx in range(1, 16):  # Columns A to O
    cell = revenue_sheet.cell(header_row, col_idx)
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    print(f"  {col_letter}: {cell.value}")

# Print first student row (row 4)
print(f"\nRow 4 (First Student - Gaju):")
for col_idx in range(1, 16):
    cell = revenue_sheet.cell(4, col_idx)
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    print(f"  {col_letter}: {cell.value}")

# Print totals row (row 11)
print(f"\nRow 11 (Total Fee Received):")
for col_idx in range(1, 16):
    cell = revenue_sheet.cell(11, col_idx)
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    print(f"  {col_letter}: {cell.value}")

print("\n" + "=" * 60)
print("CASHFLOW SHEET - Beginning Cash Row:")
print("=" * 60)

cashflow_sheet = wb["Cashflow Statement"]
# Row 3 is Beginning Cash
print("\nRow 3 (Beginning Cash):")
for col_idx in range(1, 16):
    cell = cashflow_sheet.cell(3, col_idx)
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    print(f"  {col_letter}: {cell.value}")

wb.close()
