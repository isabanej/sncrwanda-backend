import openpyxl

wb = openpyxl.load_workbook(r"c:\dev\sncrwanda-backend\frontend-new\public\SNC Rwanda - 12 Month Income & Expense Projection = CashflowProjection.xlsx", data_only=True)

revenue_sheet = wb["School Fees received - Revenue"]

print("TOTAL FEES BY MONTH (Row 11):")
print("=" * 60)

months = ["Jan", "Fab", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sept", "Oct", "Nov", "Dec"]
for i, month in enumerate(months):
    col_idx = 3 + i  # Column C = 3, D = 4, etc.
    col_letter = openpyxl.utils.get_column_letter(col_idx)
    cell_value = revenue_sheet.cell(11, col_idx).value
    if cell_value and cell_value != 0:
        print(f"{month:8} (Col {col_letter}): {cell_value:>12,.0f}")
    else:
        print(f"{month:8} (Col {col_letter}): {0:>12,}")

print("\n" + "=" * 60)
print("STUDENT DETAILS:")
print("=" * 60)

for row_idx in range(4, 11):
    name_cell = revenue_sheet.cell(row_idx, 1)
    if name_cell.value and isinstance(name_cell.value, str):
        name = name_cell.value
        print(f"\n{name}:")
        for i, month in enumerate(months):
            col_idx = 3 + i
            amount = revenue_sheet.cell(row_idx, col_idx).value
            if amount and amount != 0:
                print(f"  {month}: {amount:,.0f}")

wb.close()
