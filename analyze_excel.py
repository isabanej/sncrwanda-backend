import openpyxl
import sys

file_path = r"frontend-new\public\SNC Rwanda - 12 Month Income & Expense Projection = CashflowProjection.xlsx"

print("=" * 80)
print("EXCEL FILE STRUCTURE ANALYSIS")
print("=" * 80)

wb = openpyxl.load_workbook(file_path)

print(f"\nSheets in workbook: {wb.sheetnames}")

# Analyze Cashflow Statement sheet
print("\n" + "=" * 80)
print("CASHFLOW STATEMENT SHEET")
print("=" * 80)

sheet = wb['Cashflow Statement']
print(f"\nTotal rows: {sheet.max_row}")
print(f"Total columns: {sheet.max_column}")

print("\nFirst 3 rows, first 15 columns:")
for row_idx in range(1, 4):
    row_data = []
    for col_idx in range(1, 16):
        cell = sheet.cell(row_idx, col_idx)
        value = str(cell.value)[:20] if cell.value else "None"
        row_data.append(f"{value:20}")
    print(f"Row {row_idx}: {' | '.join(row_data)}")

print("\n" + "=" * 80)
print("SEARCHING FOR MONTH NAMES IN HEADER ROWS")
print("=" * 80)

months = ["January", "February", "March", "April", "May", "June", 
          "July", "August", "September", "October", "November", "December"]

for row_idx in range(1, 6):  # Check first 5 rows
    print(f"\nRow {row_idx}:")
    found_months = []
    for col_idx in range(1, sheet.max_column + 1):
        cell = sheet.cell(row_idx, col_idx)
        if cell.value and str(cell.value).strip() in months:
            found_months.append(f"Col {col_idx}={cell.value}")
    if found_months:
        print(f"  Found: {', '.join(found_months)}")
    else:
        print(f"  No month names found")

# Analyze School Fees sheet
print("\n" + "=" * 80)
print("SCHOOL FEES REVENUE SHEET")
print("=" * 80)

if 'School Fees received - Revenue' in wb.sheetnames:
    sheet2 = wb['School Fees received - Revenue']
    print(f"\nTotal rows: {sheet2.max_row}")
    print(f"Total columns: {sheet2.max_column}")
    
    print("\nFirst 3 rows, first 10 columns:")
    for row_idx in range(1, 4):
        row_data = []
        for col_idx in range(1, 11):
            cell = sheet2.cell(row_idx, col_idx)
            value = str(cell.value)[:15] if cell.value else "None"
            row_data.append(f"{value:15}")
        print(f"Row {row_idx}: {' | '.join(row_data)}")

print("\n" + "=" * 80)
