import openpyxl

wb = openpyxl.load_workbook(r"c:\dev\sncrwanda-backend\frontend-new\public\SNC Rwanda - 12 Month Income & Expense Projection = CashflowProjection.xlsx", data_only=True)

cashflow_sheet = wb["Cashflow Statement"]

print("EXPENSE CATEGORIES:")
print("=" * 80)

months = ["Jan", "Fab", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sept", "Oct", "Nov", "Dec"]
month_indices = list(range(3, 15))  # Columns C to N

# Check rows 10-31 for expenses
for row_idx in range(10, 32):
    # Column B has category name
    category_cell = cashflow_sheet.cell(row_idx, 2)
    if category_cell.value:
        category = str(category_cell.value).strip()
        has_data = False
        expense_data = {}
        
        for i, col_idx in enumerate(month_indices):
            cell_value = cashflow_sheet.cell(row_idx, col_idx).value
            if cell_value and cell_value != 0:
                has_data = True
                expense_data[months[i]] = cell_value
        
        if has_data:
            print(f"\n{category}:")
            for month, amount in expense_data.items():
                print(f"  {month:8}: {amount:>12,.0f}")

wb.close()
