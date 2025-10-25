import openpyxl

wb = openpyxl.load_workbook('frontend-new/public/SNC Rwanda - 12 Month Income & Expense Projection = CashflowProjection.xlsx')
sheet = wb['School Fees received - Revenue']

print("=" * 80)
print("REVENUE SHEET DETAILED ANALYSIS")
print("=" * 80)

# Show header
print("\nHeader Row (row 3):")
for col_idx in range(1, 16):
    cell_value = sheet.cell(3, col_idx).value
    print(f"  Col {col_idx:2d}: {cell_value}")

print("\n" + "=" * 80)
print("DATA ROWS:")
print("=" * 80)

for row_idx in range(4, 15):
    student_name = sheet.cell(row_idx, 1).value
    fee_type = sheet.cell(row_idx, 2).value
    
    if not student_name:
        continue
    
    print(f"\nRow {row_idx}: Student='{student_name}', Type='{fee_type}'")
    
    # Check all month columns
    for col_idx in range(3, 15):  # Columns C-N (months Jan-Dec)
        amount = sheet.cell(row_idx, col_idx).value
        if amount and isinstance(amount, (int, float)) and amount > 0:
            month_name = sheet.cell(3, col_idx).value
            print(f"  {month_name}: ${amount:,.2f}")

wb.close()
