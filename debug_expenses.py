#!/usr/bin/env python3
import openpyxl
from decimal import Decimal

# Load workbook
file_path = r"c:\dev\sncrwanda-backend\frontend-new\public\SNC Rwanda - 12 Month Income & Expense Projection = CashflowProjection.xlsx"
workbook = openpyxl.load_workbook(file_path, data_only=True)

# Get Cashflow Statement sheet
sheet = workbook["Cashflow Statement"]

print("=" * 80)
print("EXPENSE DATA ANALYSIS")
print("=" * 80)

# Check rows 10-29, column B for categories
print("\nScanning rows 10-29, column B for expense categories:")
for row_idx in range(10, 30):
    cell_b = sheet.cell(row_idx, 2)  # Column B (1-based)
    if cell_b.value:
        category = str(cell_b.value).strip()
        print(f"\nRow {row_idx}: {category}")
        
        # Check if there's numeric data in month columns (D-O = columns 4-15)
        has_data = False
        month_data = []
        for col_idx in range(4, 16):  # Columns D to O
            cell = sheet.cell(row_idx, col_idx)
            value = cell.value
            if value and (isinstance(value, (int, float)) or (isinstance(value, str) and value.strip())):
                try:
                    amt = float(value) if isinstance(value, (int, float)) else float(value.strip().replace(',', ''))
                    if amt > 0:
                        has_data = True
                        month_col = chr(64 + col_idx)  # Convert to column letter
                        month_data.append(f"{month_col}=${amt:.2f}")
                except:
                    pass
        
        if has_data:
            print(f"  Has data: {', '.join(month_data[:3])}{'...' if len(month_data) > 3 else ''}")
        else:
            print("  ⚠️ NO NUMERIC DATA FOUND")

print("\n" + "=" * 80)
print("MONTH HEADERS")
print("=" * 80)
print("\nRow 1, columns D-O (should contain month names):")
for col_idx in range(4, 16):
    cell = sheet.cell(1, col_idx)
    col_letter = chr(64 + col_idx)
    print(f"  Column {col_letter}: {cell.value}")
