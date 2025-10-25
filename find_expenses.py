import openpyxl

file_path = r"frontend-new\public\SNC Rwanda - 12 Month Income & Expense Projection = CashflowProjection.xlsx"
wb = openpyxl.load_workbook(file_path)
sheet = wb['Cashflow Statement']

print("=" * 80)
print("EXPENSE ROWS IN CASHFLOW STATEMENT")
print("=" * 80)

# Look for rows with expense data (usually after "Cash Out" or similar header)
print("\nScanning column B for category names (rows 1-50):\n")
for row_idx in range(1, 51):
    cell = sheet.cell(row_idx, 2)  # Column B
    if cell.value and isinstance(cell.value, str):
        value = str(cell.value).strip()
        if len(value) > 0 and value not in ["Cash In", "Cash Out", "None"]:
            # Check if there's a numeric value in one of the month columns
            has_data = False
            for col in range(4, 16):  # Check month columns
                month_cell = sheet.cell(row_idx, col)
                if month_cell.value and (isinstance(month_cell.value, (int, float)) or 
                                         (isinstance(month_cell.value, str) and month_cell.value.startswith('='))):
                    has_data = True
                    break
            if has_data:
                print(f"Row {row_idx}: {value}")

print("\n" + "=" * 80)
